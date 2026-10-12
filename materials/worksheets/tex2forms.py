#!/usr/bin/env python3
"""LuaLaTeX 版ワークシートの原稿（sessions/NN.tex）から Microsoft Forms 版を作る。

原稿は LaTeX 版・Word 版と共通。各回のワークシートを「質問の並び」に読み替え、
次の2つを forms/ に書き出す。

  forms/forms-NN.md    フォームの設計書。セクション、質問、種類（長い回答・選択肢など）、
                       サブタイトルに入れる説明文を、Forms で作るときの順に並べたもの。
  forms/forms-NN.docx  Forms の「クイック インポート」に読み込ませる Word ファイル。
                       番号つきの質問と、A. B. の選択肢だけの簡素な体裁にしてある
                       （表・見出し・説明文は取り込まれないため、説明文は質問文の末尾に
                       ※で添えてある）。取り込み後の手直しは forms/README.md を参照。

記入欄の読み替え方
  1行の記入欄 \\fillin                 → テキスト（短い回答）
  記入欄 \\answerbox・罫線 \\writelines → テキスト（長い回答）
  見出しつきの表（左の列が項目名）     → 項目ごとにテキスト（長い回答）
  空欄だけの表（候補を並べて書く表）   → 表ごとに1つのテキスト（長い回答）。1件1行で、列の順に書く
  チェック項目・次回までに・注意の□   → 選択肢（複数選択）
  方眼（スケッチ）                     → ファイルのアップロード
  AI活用記録                           → 4項目のテキスト（長い回答）

  python3 tex2forms.py            全回分を forms/ に書き出す
  python3 tex2forms.py 03         第3回だけ
"""
import argparse
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from tex2docx import (AIMEMO, PERSONA, PHASES, inline, items_of, parse_blocks,  # noqa: E402
                      read_args, split_top, strip_comments)

OUT = HERE / "forms"
COURSE = "基礎演習Ⅱ　2026年度後期"
LIST_HINT = "1件につき1行、この順に書く"


# ------------------------------------------------------------
# データ構造
# ------------------------------------------------------------
@dataclass
class Question:
    title: str
    kind: str                      # short / long / multi / file
    hint: str = ""                 # Forms のサブタイトルに入れる文
    options: list = field(default_factory=list)
    required: bool = False


@dataclass
class Section:
    title: str
    desc: list = field(default_factory=list)   # セクションの説明（行の列）
    questions: list = field(default_factory=list)


@dataclass
class Form:
    title: str
    desc: list = field(default_factory=list)
    sections: list = field(default_factory=list)


KIND_LABEL = {
    "short": "テキスト（短い回答）",
    "long": "テキスト（長い回答）",
    "multi": "選択肢（複数選択可）",
    "file": "ファイルのアップロード",
}


# ------------------------------------------------------------
# 文字列（LaTeX のインライン記法 → プレーンテキスト）
# ------------------------------------------------------------
def text_of(s, nl="\n"):
    """インライン記法を解いてテキストにする。\\newline は nl、下線の空欄は ＿＿。"""
    out = []
    for text, st in inline(s):
        if text == "\n" and st is None:
            out.append(nl)
        elif st.get("underline") and not text.strip("　 "):
            out.append("＿＿")
        else:
            out.append(text)
    t = "".join(out)
    t = re.sub(r"[ \u3000]*\n[ \u3000]*", nl, t)
    t = re.sub(r"[ \u3000]*\u3000[ \u3000]*", "\u3000", t)   # \hspace の連続は1つに
    t = re.sub(r" {2,}", " ", t)
    return t.strip(" \u3000\n")


PARTICLE_END = tuple("がをにでとのはへてる・／")
PARTICLE_START = tuple("がをにでとのはへ")
SYMBOLS = re.compile(r"^[○×△＋／−\-\s高中低ＡＢＣ]+$")


def split_label(raw):
    """表の項目名を（見出し, 補足）に分ける。
    \\newline で折り返しただけのもの（「先行研究／でわかって／いること」）はつなぎ、
    補足（「背景／なぜこの課題か」「感情／＋−」）は分けて返す。"""
    parts = [p for p in text_of(raw, nl="\n").split("\n") if p]
    if len(parts) <= 1:
        return (parts[0] if parts else ""), ""
    head, rest = parts[0], parts[1:]
    sub = "".join(rest)
    small = re.search(r"\\(scriptsize|footnotesize|small)", raw) is not None
    if sub.startswith("（") or small or SYMBOLS.match(sub):
        return head, sub.strip("（）")
    if head.endswith(PARTICLE_END) or sub.startswith(PARTICLE_START):
        return head + sub, ""
    if len(head) <= 2 and len(sub) <= 2:
        return head + sub, ""
    return head, sub


def join_sentences(parts):
    """文を「。」でつなぐ（末尾の「。」の重複を避ける）"""
    return "。".join(x.rstrip("。") for x in parts if x)


def with_sub(head, sub):
    return "%s（%s）" % (head, sub) if sub else head


def blank(cell):
    return not text_of(cell).strip()


# ------------------------------------------------------------
# 変換
# ------------------------------------------------------------
class Builder:
    def __init__(self):
        self.form = None
        self.work = ""          # いまの \work の題（記入欄に名前がないときの質問名）
        self.pending = []       # 次の質問かセクションに添える説明文
        self.started = False    # 最初の \work を過ぎたか

    # --- セクション・質問 ---
    def section(self, title):
        if self.pending and self.form.sections and self.cur.questions:
            self.cur.desc.extend(self.pending)
            self.pending = []
        s = Section(title)
        self.form.sections.append(s)
        return s

    @property
    def cur(self):
        return self.form.sections[-1]

    def add(self, q):
        """質問を追加する。溜まっていた説明文は、セクション最初の質問ならセクションの
        説明に、そうでなければ質問のサブタイトルに入れる。"""
        if self.pending:
            if not self.cur.questions:
                self.cur.desc.extend(self.pending)
            else:
                q.hint = join_sentences([*self.pending, q.hint])
            self.pending = []
        self.cur.questions.append(q)

    def note(self, text):
        text = text.strip()
        if text:
            if self.started:
                self.pending.append(text)
            else:
                self.form.desc.append(text)

    # --- 表 ---
    def table(self, spec, body, header):
        rows = [r for r in split_top(body, r"\\") if r.strip()]
        rows = [[re.sub(r"^\s*\\SetRow\{[^}]*\}", "", c).strip() for c in split_top(r.strip(), "&")]
                for r in rows]
        hdr = [with_sub(*split_label(c)) for c in rows[0]] if header else None
        data = rows[1:] if header else rows
        ncol = max(len(r) for r in rows)
        # 全行に同じ文字が入っている列（「分」など）は単位扱いで見出しに折り込む
        unit = set()
        for c in range(ncol):
            vals = [text_of(r[c]) if c < len(r) else "" for r in data]
            if len(data) > 1 and all(vals) and len(set(vals)) == 1:
                unit.add(c)
                if hdr:
                    hdr[c] = "%s（%s）" % (hdr[c], vals[0])
        group = []   # 空欄だけの行（候補を並べる行）をまとめる
        labeled = False   # 項目名つきの行が先にあったか

        def flush():
            if not group:
                return
            cols = sorted({c for r in group for c in r})
            names = []
            for c in cols:
                h = hdr[c] if hdr and c < len(hdr) else ""
                if h and h not in names:
                    names.append(h)
            title = "／".join(names) or self.work
            if labeled:
                title = "その他：" + title
            self.add(Question(title, "long", LIST_HINT))
            group.clear()

        for r in data:
            cells = list(r) + [""] * (ncol - len(r))
            inputs = [c for c in range(ncol) if blank(cells[c]) or c in unit]
            labels = [(c, cells[c]) for c in range(ncol) if c not in inputs]
            texts = [(c, raw) for c, raw in labels if not re.fullmatch(r"[\d０-９]+", text_of(raw))]
            if not inputs:   # 空欄のない行：記入例か、穴埋めの型
                flush()
                t = [text_of(raw, nl=" ") for _, raw in labels]
                if any("＿" in x or "①" in x for x in t):
                    self.add(Question(t[0], "long", "　".join(t[1:])))
                else:
                    self.note("　".join(t))
                continue
            if not texts:    # 番号だけ、または何もない行 → まとめて1問
                group.append(inputs)
                continue
            flush()
            labeled = True
            heads, subs = [], []
            for _, raw in texts:
                h, s = split_label(raw)
                heads.append(h)
                if s:
                    subs.append(s)
            title = "　".join(heads)
            col_hint = ""
            if hdr:
                names = [hdr[c] for c in inputs if c < len(hdr) and hdr[c]]
                if len(names) == 1:
                    col_hint = names[0]
                elif names:
                    col_hint = "次の順に書く：" + "／".join(names)
            hint = join_sentences(["、".join(subs), col_hint])
            self.add(Question(title, "long", hint))
        flush()

    # --- □ の並び ---
    def checks(self, title, items, hint=""):
        self.add(Question(title, "multi", hint, options=items))

    # --- ブロック列 ---
    def blocks(self, nodes):
        i = 0
        while i < len(nodes):
            n = nodes[i]
            if n[0] == "para":
                t = text_of(n[1], nl=" ")
                if "＿" in t:
                    head, _, rest = t.partition("：")
                    self.add(Question(head, "short", rest) if rest else Question(t, "short"))
                else:
                    self.note(t)
            elif n[0] == "cmd":
                self.cmd(n[1], n[2], n[3])
            else:
                self.env(n[1], n[2], n[3], n[4])
            i += 1

    def cmd(self, name, opts, args):
        if name == "session":
            num, ph, date, title, tag, group = args
            self.form = Form("第%s回　%s" % (num, text_of(title)))
            head = "%s　第%s回　%s｜%s" % (COURSE, num, date, PHASES[ph][2])
            if tag.strip():
                head += "｜" + text_of(tag)
            self.form.desc.append(head)
            s = self.section("回答者")
            s.questions.append(Question("学籍番号", "short", required=True))
            s.questions.append(Question("氏名", "short", required=True))
            if group.strip():
                s.questions.append(Question("グループ", "short"))
        elif name == "goals":
            self.form.desc.append("この回のねらい")
            self.form.desc.extend("・" + text_of(it, nl=" ") for it in items_of(args[0]))
        elif name == "work":
            self.started = True
            self.work = text_of(args[1], nl=" ")
            self.section("%s　%s" % (text_of(args[0]), self.work))
        elif name == "instr":
            self.note(text_of(args[0], nl=" "))
        elif name == "answerbox":
            self.add(Question(text_of(opts[0], nl=" ") if opts[0] else self.work, "long"))
        elif name == "twoboxes":
            self.add(Question(text_of(args[0], nl=" "), "long"))
            self.add(Question(text_of(args[1], nl=" "), "long"))
        elif name == "gridbox":
            self.add(Question(text_of(opts[0], nl=" ") if opts[0] else self.work, "file",
                              "手描きのスケッチを撮影した画像をアップロードする"))
        elif name == "writelines":
            self.add(Question(self.work, "long"))
        elif name == "fillin":
            t = text_of(args[0], nl=" ")
            self.add(Question(t, "short"))
        elif name == "aimemo":
            self.section("AI活用記録（この回の分）")
            self.work = "AI活用記録"
            self.blocks(parse_blocks(AIMEMO))
        elif name == "personacard":
            rows = [r for r in split_top(parse_blocks(PERSONA)[0][4], r"\\") if r.strip()]
            names = []
            for r in rows[1:]:
                raw = re.sub(r"^\s*\\SetRow\{[^}]*\}", "", split_top(r.strip(), "&")[0])
                names.append(with_sub(*split_label(raw)))
            self.add(Question("ペルソナ%s" % args[0], "long",
                              "名前／" + "／".join(names) + " の順に、1行ずつ書く"))
        else:
            raise ValueError("未対応のコマンド: \\" + name)

    def env(self, name, opts, args, body):
        if name == "minipage":
            self.blocks(parse_blocks(body))
        elif name == "promptbox":
            self.note("プロンプト例「%s」：%s" % (text_of(opts[0] or ""), text_of(body, nl=" ")))
        elif name == "caution":
            title = text_of(args[0])
            parts = [p.strip() for p in text_of(body, nl=" ").split("□")]
            if len(parts) > 1:
                self.checks(title, [p for p in parts[1:] if p], parts[0])
            else:
                self.note("%s：%s" % (title, parts[0]))
        elif name == "checks":
            self.checks(self.work, [text_of(it, nl=" ") for it in items_of(body)])
        elif name == "nexttime":
            self.section("次回までに")
            self.work = "次回までに"
            self.checks("次回までに取り組むこと（できたものにチェック）",
                        [text_of(it, nl=" ") for it in items_of(body)])
        elif name in ("wstab", "wstabh"):
            self.table(args[0], body, name == "wstabh")
        else:
            raise ValueError("未対応の環境: " + name)


def convert(num):
    src = strip_comments((HERE / "sessions" / ("%s.tex" % num)).read_text(encoding="utf-8"))
    b = Builder()
    b.blocks(parse_blocks(src))
    if b.pending:                      # 末尾に残った説明文は最後の質問に添える
        q = b.cur.questions[-1] if b.cur.questions else None
        if q:
            q.hint = join_sentences([*b.pending, q.hint])
    form = b.form
    form.sections = [s for s in form.sections if s.questions or s.desc]
    return form


# ------------------------------------------------------------
# 書き出し
# ------------------------------------------------------------
def write_md(form, num, path):
    L = ["# %s（Microsoft Forms 版）" % form.title, "",
         "sessions/%s.tex から tex2forms.py で生成した設計書。Forms で作るとき、"
         "または forms-%s.docx を取り込んだあとの確認に使う。" % (num, num), "",
         "## フォームの設定", "",
         "- タイトル：%s" % form.title,
         "- 説明："]
    L += ["  %s" % d for d in form.desc]
    L += ["- 回答できるユーザー：所属組織内のユーザーのみ（記名で回答を記録）", ""]
    k = 0
    for si, s in enumerate(form.sections, 1):
        L += ["## セクション %d　%s" % (si, s.title), ""]
        if s.desc:
            L += ["セクションの説明：", ""]
            L += ["> %s" % d for d in s.desc]
            L += [""]
        for q in s.questions:
            k += 1
            tag = KIND_LABEL[q.kind] + ("・必須" if q.required else "")
            L.append("%d. **%s**　%s" % (k, q.title, tag))
            if q.hint:
                L.append("   - サブタイトル：%s" % q.hint)
            for o in q.options:
                L.append("   - □ %s" % o)
        L.append("")
    path.write_text("\n".join(L).rstrip() + "\n", encoding="utf-8")


def write_docx(form, path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "游ゴシック"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "游ゴシック")
    doc.core_properties.title = form.title
    doc.add_paragraph(form.title, style="Title")
    for d in form.desc:
        doc.add_paragraph(d)
    doc.add_paragraph()
    k = 0
    for s in form.sections:
        for qi, q in enumerate(s.questions):
            k += 1
            line = q.title
            if q.kind == "file":
                line += "（%s）" % (q.hint or "画像をアップロード")
            elif q.hint:
                line += "（%s）" % q.hint
            if qi == 0 and s.title != "回答者":
                line = "【%s】%s" % (s.title, line)
                if s.desc:
                    line += "　※" + "　".join(s.desc)
            if q.required:
                line += "（必須）"
            doc.add_paragraph("%d. %s" % (k, line))
            for oi, o in enumerate(q.options):
                doc.add_paragraph("%s. %s" % (chr(ord("A") + oi), o))
            doc.add_paragraph()
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("nums", nargs="*", help="回番号（2桁）。省略すると全回")
    ap.add_argument("-o", "--outdir", type=Path, default=OUT)
    a = ap.parse_args()
    nums = a.nums or sorted(p.stem for p in (HERE / "sessions").glob("*.tex"))
    a.outdir.mkdir(parents=True, exist_ok=True)
    for num in nums:
        form = convert(num)
        write_md(form, num, a.outdir / ("forms-%s.md" % num))
        write_docx(form, a.outdir / ("forms-%s.docx" % num))
        nq = sum(len(s.questions) for s in form.sections)
        print("%s: %s  %dセクション %d問" % (num, form.title, len(form.sections), nq))


if __name__ == "__main__":
    main()
